"""
Combined System: Masked-Face Recognition & PPE Detection — PPT
Merges the content of generate_facerecog_pptx.py and generate_ppe_pptx.py into a single, cohesive presentation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

TITLE   = prs.slide_layouts[0]
CONTENT = prs.slide_layouts[1]
SECTION = prs.slide_layouts[2] # Section Header
BLANK   = prs.slide_layouts[6]

# ── Path to training-run images ──────────────────────────────────────
RUNS_BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         "runs", "detect", "runs", "detect")

# ─────────────────────────────────────────────────────────────────────
# Helper functions
# ─────────────────────────────────────────────────────────────────────

def title_slide(title, subtitle):
    s = prs.slides.add_slide(TITLE)
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle

def section_slide(title, subtitle=""):
    s = prs.slides.add_slide(SECTION)
    s.shapes.title.text = title
    if subtitle and len(s.placeholders) > 1:
        s.placeholders[1].text = subtitle

def bullet_slide(title, bullets, levels=None):
    s = prs.slides.add_slide(CONTENT)
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = bullets[0]
    if levels and 0 in levels:
        tf.paragraphs[0].level = levels[0]
    for i, b in enumerate(bullets[1:], 1):
        p = tf.add_paragraph()
        p.text = b
        p.level = levels.get(i, 0) if levels else 0

def table_slide(title, headers, rows, col_widths=None, footer=None):
    s = prs.slides.add_slide(BLANK)
    tx = s.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.7))
    p = tx.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True

    nr = len(rows) + 1
    nc = len(headers)
    shape = s.shapes.add_table(nr, nc, Inches(0.8), Inches(1.2),
                               Inches(11.7), Inches(min(0.45 * nr, 5.5)))
    tbl = shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Inches(w)
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = h
        for pr in c.text_frame.paragraphs:
            pr.font.size = Pt(12); pr.font.bold = True
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.text = v
            for pr in c.text_frame.paragraphs:
                pr.font.size = Pt(11)
    if footer:
        ft = s.shapes.add_textbox(Inches(0.8), Inches(1.2 + 0.45 * nr + 0.15),
                                  Inches(11), Inches(0.4))
        fp = ft.text_frame.paragraphs[0]
        fp.text = footer
        fp.font.size = Pt(12)
        fp.font.italic = True
    return s

def table_and_bullets_slide(title, headers, rows, col_widths, bullets, sub_heading="Interim System Metrics"):
    s = prs.slides.add_slide(BLANK)
    tx = s.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.7))
    p = tx.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True

    nr = len(rows) + 1
    nc = len(headers)
    tbl_height = 0.42 * nr
    shape = s.shapes.add_table(nr, nc, Inches(0.8), Inches(1.2),
                               Inches(11.7), Inches(tbl_height))
    tbl = shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Inches(w)
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = h
        for pr in c.text_frame.paragraphs:
            pr.font.size = Pt(12); pr.font.bold = True
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.text = v
            for pr in c.text_frame.paragraphs:
                pr.font.size = Pt(11)

    bullet_top = 1.2 + tbl_height + 0.3
    bx = s.shapes.add_textbox(Inches(0.8), Inches(bullet_top), Inches(11), Inches(7.0 - bullet_top))
    tf = bx.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = sub_heading
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    for b in bullets:
        bp = tf.add_paragraph()
        bp.text = b
        bp.font.size = Pt(12)
    return s

def image_slide(title, img_path):
    s = prs.slides.add_slide(BLANK)
    txBox = s.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.7))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    if os.path.exists(img_path):
        s.shapes.add_picture(img_path, Inches(1.0), Inches(1.2), Inches(11.3), Inches(5.8))
    else:
        t2 = s.shapes.add_textbox(Inches(2), Inches(3), Inches(9), Inches(1))
        t2.text_frame.paragraphs[0].text = f"[Image not found: {img_path}]"
    return s

def two_image_slide(title, img1, img2, cap1="", cap2=""):
    s = prs.slides.add_slide(BLANK)
    txBox = s.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11), Inches(0.6))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    img_top, img_h, img_w = Inches(1.0), Inches(5.2), Inches(5.8)
    if os.path.exists(img1):
        s.shapes.add_picture(img1, Inches(0.5), img_top, img_w, img_h)
    if os.path.exists(img2):
        s.shapes.add_picture(img2, Inches(6.8), img_top, img_w, img_h)
    if cap1:
        c1 = s.shapes.add_textbox(Inches(0.5), Inches(6.3), img_w, Inches(0.4))
        c1.text_frame.paragraphs[0].text = cap1
        c1.text_frame.paragraphs[0].font.size = Pt(12)
        c1.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    if cap2:
        c2 = s.shapes.add_textbox(Inches(6.8), Inches(6.3), img_w, Inches(0.4))
        c2.text_frame.paragraphs[0].text = cap2
        c2.text_frame.paragraphs[0].font.size = Pt(12)
        c2.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return s

# =====================================================================
# PAGE 1 — Title
# =====================================================================
title_slide(
    "Automated Staff Entry & Compliance System",
    "Progress Report -- Proof of Concept\n\nIntegrating Masked-Face Recognition & PPE Detection"
)

# =====================================================================
# PAGE 2 — 1. Executive Summary
# =====================================================================
bullet_slide("1. Executive Summary", [
    "Dual-engine entry-control system ensuring both Identity and Safety Compliance.",
    "System 1: Masked-Face Recognition — Identifies staff from partial/occluded faces (surgical mask, hair cap) using RetinaFace + FaceNet.",
    "System 2: PPE Detection Engine — Validates mandated medical protective equipment (mask, gloves, goggles, coverall) using YOLOv11.",
    "Current phase delivers working proofs-of-concept for both pipelines, demonstrating high accuracy and functional completeness.",
    "Approved scope implemented and executable end-to-end: ~60% (across both systems).",
    "Remaining work: Integration into a single unified deployment, field testing, and resolving concurrent tracking limitations.",
])


# =====================================================================
# SECTION: Masked-Face Staff Entry Recognition
# =====================================================================
section_slide("Part 1: Masked-Face Staff Entry Recognition", "Identity Verification under Occlusion")

# PAGE 3
table_slide(
    "2. Implementation — Modules Delivered (Face Recognition)",
    ["Module", "File", "Function"],
    [
        ["Detection + Embedding + Matching", "face_engine.py", "Core recognition logic: RetinaFace detection, FaceNet embedding, cosine matching"],
        ["Storage layer", "db.py, schema.sql", "SQLite persistence for staff, embeddings, and entry logs (WAL mode)"],
        ["Configuration", "config.py", "Centralized tunables: thresholds, camera settings, capture parameters"],
        ["Enrollment capture", "capture_utils.py", "Standalone webcam burst-capture window (8 photos), subprocess"],
        ["Live entry monitor", "entry_monitor.py", "Real-time camera loop: detect, embed, match, draw decision, debounced logging"],
        ["Admin dashboard", "Home.py", "Streamlit entrypoint, model warm-up, summary metrics"],
        ["Staff registration UI", "pages/1_Register_Staff.py", "Guided burst-capture enrollment flow"],
        ["Authorization & Logs", "pages/2 & 3", "Manage staff authorization, filterable log table with snapshot preview"],
    ],
    col_widths=[2.5, 2.8, 6.4],
)

# PAGE 4
bullet_slide("2.2 System Architecture (Face Recognition)", [
    "Webcam burst capture (enrollment) or live camera feed (entry)",
    "Face detection -- RetinaFace, occlusion-tuned confidence threshold (0.6)",
    "Embedding generation -- FaceNet InceptionResnetV1, 512-d, L2-normalized",
    "Open-set matching -- cosine similarity, per-person maximum",
    "Decision (AUTHORIZED / DENIED / UNKNOWN) logged to SQLite with a snapshot",
    "Reviewed via the Streamlit admin dashboard",
    "",
    "Streamlit admin app and standalone entry-monitor camera loop run concurrently against the same SQLite database, coordinated via Write-Ahead Logging (WAL) mode.",
])

# PAGE 5
table_slide(
    "3. Technical Accuracy and Use of Best Practices (Face)",
    ["Decision", "Rationale"],
    [
        ["RetinaFace for detection", "Selected specifically for robustness to facial occlusion (mask, cap), rather than a generic detector."],
        ["Detector confidence threshold lowered to 0.6", "Occluded faces consistently scored below the typical 0.9 default, causing missed detections."],
        ["FaceNet (InceptionResnetV1)", "Highest-performing embedding model from baseline research (97.7% training accuracy)."],
        ["Open-set cosine-similarity matching", "Allows new staff to be enrolled without retraining any model, unlike a closed-set classifier."],
        ["Background-thread detection worker", "Keeps the live video preview rendering at camera frame rate while the slower detect+embed pass runs asynchronously."],
    ],
    col_widths=[4.5, 7.2],
)

# PAGE 6
s = prs.slides.add_slide(BLANK)
tx = s.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.7))
p = tx.text_frame.paragraphs[0]
p.text = "4.1 Interim Results -- Baseline (Prior Research - Face)"
p.font.size = Pt(28)
p.font.bold = True
bx = s.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11), Inches(1.2))
tf = bx.text_frame
tf.word_wrap = True
tf.paragraphs[0].text = "Five model architectures evaluated for partial-face recognition on a custom dataset (423 original images, augmented to 2,202 via rotation, shear, zoom, horizontal flip, and brightness adjustments)."
tf.paragraphs[0].font.size = Pt(14)
bp = tf.add_paragraph()
bp.text = "FaceNet achieved the strongest result and was carried forward into this implementation."
bp.font.size = Pt(14)
shape = s.shapes.add_table(6, 2, Inches(0.8), Inches(2.7), Inches(7.0), Inches(3.0))
tbl = shape.table
tbl.columns[0].width = Inches(3.0)
tbl.columns[1].width = Inches(4.0)
data = [
    ["Metric", "Value"],
    ["Original dataset size", "423 images"],
    ["Augmented dataset size", "2,202 images"],
    ["Best-performing model", "FaceNet (InceptionResnetV1)"],
    ["Training accuracy", "97.7%"],
    ["Validation accuracy", "85.5%"],
]
for i, row in enumerate(data):
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.text = val
        for pr in c.text_frame.paragraphs:
            pr.font.size = Pt(12)
            pr.font.bold = (i == 0)

# PAGE 7
table_and_bullets_slide(
    "4.2-4.3 Current POC Configuration & Interim Metrics (Face)",
    ["Parameter", "Value", "Status"],
    [
        ["Embedding dimension",           "512",                       "Fixed (FaceNet output)"],
        ["Detector confidence threshold",  "0.6",                      "Tuned for occlusion, empirically set"],
        ["Minimum images to enroll",       "3 successful detections",  "Configured"],
        ["Burst capture",                  "8 photos at 1.2s interval","Configured"],
    ],
    col_widths=[3.5, 3.5, 4.7],
    bullets=[
        "Entry-logs dashboard produces live, interpretable metrics for every recognition attempt: AUTHORIZED / DENIED / UNKNOWN counts, filterable by date and staff name, each with an associated snapshot image.",
        "Dedicated calibration tool computes genuine-pair vs. impostor-pair cosine-similarity statistics (min, mean, max, std) from enrolled staff embeddings.",
        "Scheduled as the next data-collection milestone once a representative set of staff has been enrolled.",
    ],
)

# PAGE 8
bullet_slide("5. Problem-Solving & Technical Refinement (Face)", [
    "Occluded faces missed by the detector",
    "Problem: Default RetinaFace confidence threshold (0.9) rejected masked/capped faces.",
    "Fix: Lowered detector confidence threshold to 0.6, restoring detection recall.",
    "",
    "Streamlit's built-in camera input unsuitable for burst enrollment",
    "Problem: st.camera_input captures only a single shot per rerun, too limited for an eight-photo enrollment burst.",
    "Fix: Built standalone OpenCV subprocess (capture_utils.py) with live preview and countdown.",
    "",
    "Live video preview freezing during recognition",
    "Problem: Synchronous detect+embed on every frame would block the video feed (CPU-bound, slow).",
    "Fix: Moved detection/embedding to a background thread; main loop renders latest completed result.",
], levels={0: 0, 1: 1, 2: 1, 4: 0, 5: 1, 6: 1, 8: 0, 9: 1, 10: 1})


# =====================================================================
# SECTION: Medical PPE Detection Engine
# =====================================================================
section_slide("Part 2: Medical PPE Detection Engine", "Safety Compliance Verification")

# PAGE 9
table_slide(
    "2. Implementation — Modules Delivered (PPE)",
    ["Module", "File", "Function"],
    [
        ["Dataset acquisition", "step1_download.py", "Automated download from Kaggle API and Roboflow API"],
        ["Dataset merging", "step2_merge.py", "Unifies datasets: VOC→YOLO conversion, class mapping, deduplication"],
        ["Training (merged)", "step3_train.py", "YOLOv11n on merged 3-class dataset (mask, glove, cap)"],
        ["Training (medical)", "step3_train_medical_ppe.py", "YOLOv11n on dedicated 4-class dataset"],
        ["Training (augmented) ★", "step3_train_roboflow.py", "YOLOv11n on augmented dataset (mAP@50 = 97.5%)"],
        ["Compliance arbiter", "step4_compliance_arbiter.py", "Real-time webcam inference with temporal state machine"],
        ["Colab training", "train_colab.ipynb", "Google Colab GPU training notebook"],
    ],
    col_widths=[2.3, 3.0, 6.4],
)

# PAGE 10
bullet_slide("2.2 System Architecture (PPE)", [
    "Webcam captures live video feed at entry point",
    "YOLOv11 inference — detects PPE items (mask, gloves, goggles, coverall)",
    "Detection parsing — extracts detected class names from YOLO output",
    "Temporal history buffer — rolling deque of last 15–100 frames' detections",
    "Temporal validation — a class is 'validated' only when detected in ≥ 5 frames within the buffer",
    "State machine — three states: IDLE → EVALUATING (10s window) → GRANTED",
    "",
    "On compliance success: door-open signal held for 5 seconds. On timeout: set-difference identifies which PPE items are missing and sends dispense signals.",
])

# PAGE 11
table_slide(
    "3. Technical Accuracy and Use of Best Practices (PPE)",
    ["Decision", "Rationale"],
    [
        ["YOLOv11 Nano (yolo11n.pt) for detection", "2.6M parameters, fits in 4 GB VRAM, real-time inference (~30+ FPS at 416×416), COCO-pretrained."],
        ["Image size 416 instead of 640", "~2.5× faster training per epoch; sufficient resolution for PPE items; negligible accuracy trade-off."],
        ["Data augmentation (blur, greyscale, noise)", "Forces the model to learn shape and texture rather than relying on colour. Proved decisive: +29% mAP improvement."],
        ["Dataset merging with class harmonisation", "15+ label variations (face_mask, with_mask, medical_mask) mapped to 3 canonical classes."],
        ["Temporal validation (rolling history)", "Prevents single-frame false positives from triggering compliance decisions; smooths noisy detections."],
    ],
    col_widths=[4.5, 7.2],
)

# PAGE 12
s = prs.slides.add_slide(BLANK)
tx = s.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.7))
p = tx.text_frame.paragraphs[0]
p.text = "4.1 Interim Results -- Datasets and Training Runs (PPE)"
p.font.size = Pt(28)
p.font.bold = True
bx = s.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11), Inches(1.4))
tf = bx.text_frame
tf.word_wrap = True
tf.paragraphs[0].text = "Three training runs compared, all using YOLOv11 Nano. Data augmentation (blur, greyscale, noise, brightness, flip) applied via Roboflow before export."
tf.paragraphs[0].font.size = Pt(14)
bp = tf.add_paragraph()
bp.text = "The Roboflow Augmented run achieved the best results and was selected as the production model."
bp.font.size = Pt(14)
shape = s.shapes.add_table(5, 5, Inches(0.8), Inches(2.8), Inches(11.5), Inches(2.5))
tbl = shape.table
tbl.columns[0].width = Inches(2.5)
tbl.columns[1].width = Inches(1.5)
tbl.columns[2].width = Inches(2.5)
tbl.columns[3].width = Inches(2.5)
tbl.columns[4].width = Inches(2.5)
data = [
    ["Metric", "Merged PPE", "Medical PPE", "Roboflow Aug. ★"],
    ["Dataset size (train)", "9,630", "3,352", "8,304"],
    ["Classes", "3", "4", "4"],
    ["Best mAP@50", "—", "68.5%", "97.5%"],
    ["Best mAP@50-95", "—", "40.0%", "97.3%"],
]
for i, row in enumerate(data):
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.text = val
        for pr in c.text_frame.paragraphs:
            pr.font.size = Pt(12)
            pr.font.bold = (i == 0)

kx = s.shapes.add_textbox(Inches(0.8), Inches(5.6), Inches(11), Inches(0.8))
kf = kx.text_frame
kf.word_wrap = True
kf.paragraphs[0].text = "Key Insight: Augmentation alone boosted mAP@50 by +29%. The un-augmented model overfit and generalised poorly."
kf.paragraphs[0].font.size = Pt(14)
kf.paragraphs[0].font.bold = True

# PAGE 13
image_slide("4.1 Results — Training Curves (★ PPE Roboflow Augmented)", os.path.join(RUNS_BASE, "ppe_engine_roboflow", "results.png"))

# PAGE 14
two_image_slide("4.1 Results — Confusion Matrix (★ PPE Roboflow Augmented)",
    os.path.join(RUNS_BASE, "ppe_engine_roboflow", "confusion_matrix.png"),
    os.path.join(RUNS_BASE, "ppe_engine_roboflow", "confusion_matrix_normalized.png"),
    "Raw Counts", "Normalised")

# PAGE 15
two_image_slide("4.1 Results — F1 and Precision-Recall Curves (★ PPE)",
    os.path.join(RUNS_BASE, "ppe_engine_roboflow", "BoxF1_curve.png"),
    os.path.join(RUNS_BASE, "ppe_engine_roboflow", "BoxPR_curve.png"),
    "F1-Confidence Curve", "Precision-Recall Curve")

# PAGE 16
two_image_slide("4.1 Results — Validation Predictions (★ PPE)",
    os.path.join(RUNS_BASE, "ppe_engine_roboflow", "val_batch0_labels.jpg"),
    os.path.join(RUNS_BASE, "ppe_engine_roboflow", "val_batch0_pred.jpg"),
    "Ground Truth Labels", "Model Predictions")

# PAGE 17
table_and_bullets_slide(
    "4.2-4.3 Current POC Configuration & Interim Metrics (PPE)",
    ["Parameter", "Value", "Status"],
    [
        ["Base model",                 "YOLOv11 Nano (yolo11n.pt)",    "COCO-pretrained, 2.6M params"],
        ["Image size",                 "416×416",                       "Reduced from 640 for ~2.5× speed"],
        ["Batch size",                 "80",                            "Maximises RTX 3050 GPU utilisation"],
        ["Temporal history buffer",    "15–100 frames",                 "Configurable per arbiter variant"],
        ["Min. detection count",       "5 frames",                      "Class must appear in ≥5 frames to validate"],
    ],
    col_widths=[3.5, 3.0, 5.2],
    bullets=[
        "★ Best model: mAP@50 = 97.5%, mAP@50-95 = 97.3%, Precision = 94.9%, Recall = 92.8%.",
        "Final box loss 0.0198, class loss 0.0678 — near-perfect localisation and classification.",
        "Converged in 50 epochs (~41 minutes on RTX 3050); no early stopping triggered.",
    ],
)

# PAGE 18
bullet_slide("5. Problem-Solving & Technical Refinement (PPE)", [
    "False Positives from Single-Frame Glitches",
    "Problem: The YOLO model occasionally misclassified generic objects as PPE for a split second, leading to flickering detections.",
    "Fix: Implemented a rolling history buffer (temporal validation) requiring an item to appear in ≥ 5 frames before being verified.",
    "",
    "Unstable Compliance Decisions",
    "Problem: Momentary occlusions (e.g., turning the head) would instantly reset the compliance check if done naively per-frame.",
    "Fix: Built a deterministic state machine (IDLE → EVALUATING → GRANTED) with a 10s window to smooth out temporary drops in detection.",
    "",
    "Model Overfitting without Augmentation",
    "Problem: Un-augmented model trained on clear images overfit, reaching only 68.5% mAP@50.",
    "Fix: Applied blur, greyscale, noise, and exposure changes to force learning of shapes/textures (+29% mAP).",
    "",
    "Label Fragmentation Across Datasets",
    "Problem: Datasets used 15+ variations for the same items (e.g., face_mask, with_mask).",
    "Fix: Unified TARGET_MAPPING dictionary mapped all to 3 canonical classes.",
], levels={0: 0, 1: 1, 2: 1, 4: 0, 5: 1, 6: 1, 8: 0, 9: 1, 10: 1, 12: 0, 13: 1, 14: 1})


# =====================================================================
# SECTION: Combined Open Issues & Conclusion
# =====================================================================
section_slide("Remaining Scope & Open Issues", "Future Work")

# PAGE 19
bullet_slide("2.3 & 5.1 Remaining Scope & Open Issues (Combined)", [
    "Face Recognition:",
    "Calibration of match threshold against real enrolled-staff data.",
    "A single global cooldown bucket means concurrent unknown individuals may not be logged correctly.",
    "",
    "PPE Detection:",
    "No per-person PPE tracking — validates PPE presence in frame, not on a specific individual.",
    "Planned fix: integrate YOLOv11 person detection to pair PPE with the nearest person bounding box.",
    "",
    "System Integration (Long-term vision):",
    "Merge both engines into a unified pipeline: detect person, verify identity via FaceNet, verify PPE via YOLO, and grant access only if BOTH conditions are met.",
    "Edge deployment on NVIDIA Jetson Nano for offline entry-point use.",
], levels={0: 0, 1: 1, 2: 1, 4: 0, 5: 1, 6: 1, 8: 0, 9: 1, 10: 1})

# PAGE 20
title_slide("Thank You", "Questions?")

# ── Save ─────────────────────────────────────────────────────────────
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Final_Combined_Presentation.pptx")
prs.save(out)
print(f"Saved -> {out}")
